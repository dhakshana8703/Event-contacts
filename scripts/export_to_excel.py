"""
Excel Exporter for Dubai Business Events & Contacts Deliverables.
Generates a multi-tab, professionally styled .xlsx workbook from data/output/contacts.json and data/output/events.json.
"""
import os
import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "data", "output")
CONTACTS_JSON = os.path.join(OUTPUT_DIR, "contacts.json")
EVENTS_JSON = os.path.join(OUTPUT_DIR, "events.json")
EXCEL_OUTPUT = os.path.join(OUTPUT_DIR, "dubai_business_events_contacts.xlsx")


def export_to_excel(
    contacts_path: str = CONTACTS_JSON,
    events_path: str = EVENTS_JSON,
    output_path: str = EXCEL_OUTPUT
) -> str:
    """Read JSON deliverables and compile a formatted multi-sheet Excel file."""
    if not os.path.exists(contacts_path):
        raise FileNotFoundError(f"Contacts JSON not found: {contacts_path}")
    if not os.path.exists(events_path):
        raise FileNotFoundError(f"Events JSON not found: {events_path}")

    with open(contacts_path, "r", encoding="utf-8") as f:
        contacts = json.load(f)

    with open(events_path, "r", encoding="utf-8") as f:
        events = json.load(f)

    wb = openpyxl.Workbook()

    # Sheet 1: Contacts & Desks
    ws_contacts = wb.active
    ws_contacts.title = "Contacts & Desks"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    contact_headers = [
        "Event Name", "Event Dates", "Event Venue", "Official Event Website",
        "Event Category", "Contact Type", "Contact / Desk Name", "Job Title / Role",
        "Organization", "Work Email", "Phone Number", "LinkedIn Profile URL",
        "Verification Status", "Evidence Snippet", "Notes", "Source URL"
    ]
    ws_contacts.append(contact_headers)

    for c in contacts:
        ctype = c.get("contact_type", "person")
        if ctype == "person":
            label_type = "Person"
        elif ctype == "department_desk":
            label_type = "General Desk"
        else:
            label_type = "No Contacts Published"

        row = [
            c.get("event_name", ""),
            c.get("event_dates", ""),
            c.get("event_venue", ""),
            c.get("event_website", ""),
            c.get("event_category", ""),
            label_type,
            c.get("person_name", {}).get("normalized_value", "") if c.get("person_name") else "",
            c.get("role_as_written", {}).get("normalized_value", "") if c.get("role_as_written") else "",
            c.get("organisation", {}).get("normalized_value", "") if c.get("organisation") else "",
            c.get("email", {}).get("normalized_value", "") if c.get("email") else "",
            c.get("phone", {}).get("normalized_value", "") if c.get("phone") else "",
            c.get("linkedin_url", {}).get("normalized_value", "") if c.get("linkedin_url") else "",
            c.get("verification_status", ""),
            c.get("person_name", {}).get("supporting_snippet", "") if c.get("person_name") else "",
            c.get("notes", ""),
            c.get("contact_source_url", "")
        ]
        ws_contacts.append(row)

    # Style Header & Cells for Contacts
    for col_num in range(1, len(contact_headers) + 1):
        cell = ws_contacts.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws_contacts.iter_rows(min_row=2, max_row=ws_contacts.max_row, min_col=1, max_col=len(contact_headers)):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

    for col in ws_contacts.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_contacts.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    # Sheet 2: Events Overview
    ws_events = wb.create_sheet(title="Events Overview")
    event_headers = [
        "Event Name", "Dates", "Venue", "Category", "Official Website URL",
        "Website Status", "Listing Page URL", "Detail Page URL", "Retrieval Status", "Notes"
    ]
    ws_events.append(event_headers)

    for e in events:
        ws_events.append([
            e.get("event_name", {}).get("normalized_value", "") if e.get("event_name") else "",
            e.get("dates_text", {}).get("normalized_value", "") if e.get("dates_text") else "",
            e.get("venue", {}).get("normalized_value", "") if e.get("venue") else "",
            e.get("category", {}).get("normalized_value", "") if e.get("category") else "",
            e.get("event_website_url", {}).get("normalized_value", "") if e.get("event_website_url") else "",
            e.get("event_website_url", {}).get("extraction_status", "") if e.get("event_website_url") else "",
            e.get("listing_page_url", ""),
            e.get("event_page_url", ""),
            e.get("retrieval_status", ""),
            e.get("notes", "")
        ])

    for col_num in range(1, len(event_headers) + 1):
        cell = ws_events.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws_events.iter_rows(min_row=2, max_row=ws_events.max_row, min_col=1, max_col=len(event_headers)):
        for cell in row:
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

    for col in ws_events.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_events.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    # Freeze Header Rows
    ws_contacts.freeze_panes = "A2"
    ws_events.freeze_panes = "A2"

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    out = export_to_excel()
    print(f"Generated Excel deliverable: {out}")
